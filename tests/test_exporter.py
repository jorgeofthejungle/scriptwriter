"""Unit tests for the DOCX Script Exporter."""

import unittest
import tempfile
from pathlib import Path
import docx

from engine.exporter import DocxScriptExporter


class TestDocxExporter(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.sample_md = Path(self.temp_dir.name) / "test_script.md"
        self.sample_md.write_text(
            "# The Great Sighting at Ocracoke\n\n"
            "On Friday evening, Thomas was driving near Ocracoke (OH-kruh-koke). The town was quiet.\n\n"
            "He saw a white light hovering over the dunes.\n\n"
            "---\n# Sources\n- Test Source 1\n",
            encoding="utf-8"
        )

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_exports_valid_docx_with_metadata_and_appendix(self):
        exporter = DocxScriptExporter(str(self.sample_md))
        out_path = exporter.export(output_dir=self.temp_dir.name, story_id="STORY-TEST")

        self.assertTrue(out_path.exists())
        self.assertTrue(out_path.stat().st_size > 0)

        # Inspect generated docx structure
        doc = docx.Document(str(out_path))
        self.assertTrue(len(doc.paragraphs) > 0)
        self.assertTrue(len(doc.tables) > 0)

        # Check metadata table contents
        meta_table = doc.tables[0]
        cell_text = " ".join([c.text for row in meta_table.rows for c in row.cells])
        self.assertIn("STORY-TEST", cell_text)
        self.assertIn("Safe Outcome Verified", cell_text)

        # Check pronunciation guides in runs
        all_runs_text = [r.text for p in doc.paragraphs for r in p.runs]
        self.assertTrue(any("Ocracoke (OH-kruh-koke)" in t for t in all_runs_text))


if __name__ == "__main__":
    unittest.main()
